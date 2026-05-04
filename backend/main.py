"""
Perisai Document Generator — FastAPI Backend
Jalankan: python -m uvicorn main:app --port 8000
"""

import os, re, io, fitz, logging, datetime, tempfile, httpx
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, HTTPException
from fastapi.responses import StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

def normalize(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (text or "").lower().strip())

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger(__name__)

app = FastAPI(title="Perisai Doc Generator")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

BASE_DIR   = Path(__file__).parent
ASSETS_DIR = BASE_DIR / "assets"
TEMP_DIR   = BASE_DIR / "temp"
TEMP_DIR.mkdir(exist_ok=True)

TEMPLATE_PKS   = ASSETS_DIR / "TEMPLATE PKS MASTER.docx"
TEMPLATE_FORM  = ASSETS_DIR / "FORMULIR PENDAFTARAN CALON PERISAI 2026.pdf"
TEMPLATE_UJIAN = ASSETS_DIR / "TEMPLATE UJIAN.pdf"
STEMPEL_WADAH  = ASSETS_DIR / "stempel_wadah.png"
TTD_WADAH      = ASSETS_DIR / "ttd_wadah.png"
STEMPEL_CABANG = ASSETS_DIR / "stempel_cabang.png"
TTD_CABANG     = ASSETS_DIR / "ttd_cabang.png"

# Ketetapan default
DEFAULT_NAMA_WADAH     = "HERDIANA SIMBOLON"
DEFAULT_JABATAN_WADAH  = "KETUA WADAH"
DEFAULT_NAMA_CABANG    = "SAKINAH RAMZA"
DEFAULT_JABATAN_CABANG = "ARK"

# Kamus kode pos lokal untuk area Medan Maimun
# Kamus kode pos lokal untuk area rekrutmen Perisai
KODE_POS_LOKAL = {
    # KECAMATAN MEDAN MAIMUN
    "AUR": "20151",
    "HAMDAN": "20151",
    "JATI": "20152",
    "KAMPUNG BARU": "20158",
    "SEI MATI": "20159",
    "SUKA RAJA": "20159",

    # KECAMATAN MEDAN BARAT
    "GLUGUR KOTA": "20115",
    "KARANG BEROMBAK": "20117",
    "KESAWAN": "20111",
    "PULO BRAYAN KOTA": "20116",
    "SEI AGUL": "20117",
    "SILALAS": "20114",

    # KECAMATAN MEDAN DENAI
    "BINJAI": "20228",
    "DENAI": "20227",
    "MEDAN TENGGARA": "20228",
    "TEGAL SARI MANDALA I": "20226",
    "TEGAL SARI MANDALA II": "20226",
    "TEGAL SARI MANDALA III": "20226",
}
# ── Download dari URL langsung (GDrive atau URL lain) ─────────────────────────

def gdrive_file_id_from_url(url: str) -> Optional[str]:
    """Ekstrak file ID dari URL Google Drive (file, bukan folder)."""
    patterns = [
        r"/file/d/([a-zA-Z0-9_-]+)",
        r"[?&]id=([a-zA-Z0-9_-]+)",
        r"/open\?id=([a-zA-Z0-9_-]+)",
    ]
    for p in patterns:
        m = re.search(p, url)
        if m: return m.group(1)
    return None

async def download_file_from_url(url: str, dest_path: Path, label: str = "") -> bool:
    """
    Download file dari URL. Handle Google Drive (direct download + confirm token)
    dan URL biasa lainnya.
    """
    if not url or not url.startswith("http"):
        logger.warning(f"[{label}] URL tidak valid: {url}")
        return False

    # Coba deteksi apakah ini URL GDrive file
    file_id = gdrive_file_id_from_url(url)

    async with httpx.AsyncClient(follow_redirects=True, timeout=30) as client:
        # Gunakan direct download URL jika ini GDrive file
        download_url = f"https://drive.google.com/uc?export=download&id={file_id}" if file_id else url

        resp = await client.get(download_url)

        if resp.status_code != 200:
            logger.error(f"[{label}] HTTP {resp.status_code} dari {download_url}")
            return False

        content      = resp.content
        content_type = resp.headers.get("content-type", "")

        # GDrive virus-scan confirmation (file besar)
        if "text/html" in content_type and file_id:
            logger.info(f"[{label}] GDrive HTML response — mencoba confirm token...")
            html = content.decode("utf-8", errors="ignore")
            confirm = re.search(r'confirm=([0-9A-Za-z_\-]+)', html)
            if confirm:
                resp2 = await client.get(
                    f"https://drive.google.com/uc?export=download&confirm={confirm.group(1)}&id={file_id}"
                )
                if resp2.status_code == 200 and "text/html" not in resp2.headers.get("content-type", ""):
                    content = resp2.content
                else:
                    logger.error(f"[{label}] Confirm download gagal")
                    return False
            else:
                # Coba endpoint alternatif
                resp3 = await client.get(f"https://drive.google.com/uc?id={file_id}&export=download&confirm=t")
                if resp3.status_code == 200 and "text/html" not in resp3.headers.get("content-type", ""):
                    content = resp3.content
                else:
                    logger.error(f"[{label}] Semua endpoint GDrive gagal")
                    return False

        if len(content) < 100:
            logger.error(f"[{label}] Content terlalu kecil ({len(content)} bytes)")
            return False

        dest_path.write_bytes(content)
        logger.info(f"[{label}] Berhasil download {len(content):,} bytes → {dest_path.name}")
        return True

# ── Kode Pos ─────────────────────────────────────────────────────────────────

async def lookup_kodepos(kelurahan: str, kecamatan: str = "") -> str:
    if not kelurahan:
        return ""

    kelurahan_bersih = kelurahan.strip().upper()

    # 1. Prioritas Utama: Ambil dari kamus lokal (Instan & Anti-Gagal)
    if kelurahan_bersih in KODE_POS_LOKAL:
        logger.info(f"Kodepos dari lokal: {kelurahan_bersih} -> {KODE_POS_LOKAL[kelurahan_bersih]}")
        return KODE_POS_LOKAL[kelurahan_bersih]

    # 2. Cadangan: Jika kelurahan di luar daftar, gunakan API
    try:
        logger.info(f"Kodepos dari API untuk: {kelurahan_bersih}")
        async with httpx.AsyncClient(timeout=10) as client:
            resp = await client.get(
                "https://kodepos.vercel.app/search/",
                params={"q": kelurahan.strip()}
            )
            if resp.status_code == 200:
                data = resp.json()
                results = data.get("data", data) if isinstance(data, dict) else data
                if isinstance(results, list) and results:
                    if kecamatan:
                        kec_lower = re.sub(r"[^a-z\s]", "", kecamatan.lower())
                        for r in results:
                            sub = re.sub(r"[^a-z\s]", "", (r.get("subdistrict") or "").lower())
                            if kec_lower in sub or sub in kec_lower:
                                return r.get("postalcode", "")
                    return results[0].get("postalcode", "")
    except Exception as e:
        logger.error(f"Kodepos API error: {e}")

    # 3. Fallback Darurat: Jika API down dan kelurahan tidak ada di kamus
    return ""

# ── Request Models ────────────────────────────────────────────────────────────

class KeplingData(BaseModel):
    namaLengkap:      str
    namaCalonPerisai: Optional[str] = None
    noKTP:            str
    tempatTglLahir:   str
    noKPJ:            str
    kodeNamaWadah:    str
    alamat:           str
    kabupatenKota:    str
    kodePos:          Optional[str] = ""
    kelurahan:        Optional[str] = None
    kecamatan:        Optional[str] = None
    noTelp:           str
    email:            str
    pekerjaan:        str
    noNPWP:           Optional[str] = ""
    bank:             str
    noRek:            str
    namaPemilik:      str
    namaWadah:        Optional[str] = None
    jabatanWadah:     Optional[str] = None
    namaCabang:       Optional[str] = None
    jabatanCabang:    Optional[str] = None
    # URL langsung per file
    ttdUrl:           Optional[str] = None   # URL file ttd_calon
    materaiUrl:       Optional[str] = None   # URL file materai
    dokumenUrl:       Optional[str] = None   # URL file dokumen.pdf
    nomorSurat:       Optional[int] = None

class ExamData(BaseModel):
    namaLengkap: str
    skor:        float
    jawaban:     dict

# ── Helper ────────────────────────────────────────────────────────────────────

def resolve_penandatangan(data: KeplingData) -> tuple[str, str, str, str]:
    return (
        (data.namaWadah    or DEFAULT_NAMA_WADAH).upper(),
        (data.jabatanWadah or DEFAULT_JABATAN_WADAH).upper(),
        (data.namaCabang   or DEFAULT_NAMA_CABANG).upper(),
        (data.jabatanCabang or DEFAULT_JABATAN_CABANG).upper(),
    )

async def resolve_kodepos(data: KeplingData) -> str:
    if data.kodePos and data.kodePos.strip():
        return data.kodePos.strip()
    return await lookup_kodepos(data.kelurahan or "", data.kecamatan or "")

def safe_insert_img(page: fitz.Page, path: Optional[Path], rect: fitz.Rect, label: str = ""):
    if not path:
        logger.warning(f"insert_img [{label}]: path None"); return
    if not path.exists():
        logger.warning(f"insert_img [{label}]: tidak ada di {path}"); return
    if path.stat().st_size == 0:
        logger.warning(f"insert_img [{label}]: kosong (0 bytes)"); return
    try:
        page.insert_image(rect, filename=str(path))
        logger.info(f"insert_img [{label}]: OK ({path.stat().st_size:,} bytes)")
    except Exception as e:
        logger.error(f"insert_img [{label}]: error — {e}")

# ── PDF Generator: Formulir ───────────────────────────────────────────────────

def generate_formulir(
    data: KeplingData,
    ttd_path: Optional[Path],
    materai_path: Optional[Path],
    kode_pos: str,
) -> bytes:
    if not TEMPLATE_FORM.exists():
        raise FileNotFoundError("Template formulir tidak ditemukan di assets/")

    doc  = fitz.open(str(TEMPLATE_FORM))
    page = doc[0]

    PX, PY = 3, -2
    nama_display = (data.namaCalonPerisai or data.namaLengkap).upper()

    now        = datetime.datetime.now()
    dd, mm, yy = now.strftime("%d"), now.strftime("%m"), now.strftime("%y")

    nama_wadah, jabatan_wadah, nama_cabang, jabatan_cabang = resolve_penandatangan(data)

    # ── Teks data ─────────────────────────────────────────────
    fields = [
        ((156+PX, 134+PY), nama_display,                    9),
        ((156+PX, 158+PY), data.noKTP,                      9),
        ((156+PX, 183+PY), data.tempatTglLahir.upper(),     9),
        ((156+PX, 207+PY), data.noKPJ,                      9),
        ((156+PX, 232+PY), data.kodeNamaWadah.upper(),      9),
        ((156+PX, 256+PY), data.alamat.upper(),              9),
        ((224+PX, 284+PY), data.kabupatenKota.upper(),       9),
        ((418+PX, 284+PY), kode_pos,                         9),
        ((156+PX, 308+PY), data.noTelp,                     9),
        ((156+PX, 332+PY), data.email.lower(),               9),
        ((156+PX, 357+PY), data.pekerjaan.upper(),           9),
        ((156+PX, 381+PY), (data.noNPWP or "").upper(),     9),
        ((179+PX, 406+PY), data.bank.upper(),                9),
        ((302+PX, 406+PY), data.noRek,                      9),
        ((211+PX, 430+PY), data.namaPemilik.upper(),         9),
    ]
    for (x, y), teks, fs in fields:
        page.insert_text((x, y), teks, fontsize=fs, fontname="helv", color=(0,0,0))

    def insert_centered(teks: str, cx: float, y: float, fs: int):
        w = fitz.get_text_length(teks, fontname="helv", fontsize=fs)
        page.insert_text((cx - w/2, y), teks, fontsize=fs, fontname="helv", color=(0,0,0))

    # ── Area TTD (bagian bawah form) ──────────────────────────
    # Tempat TTD (kiri bawah, baris "Medan, dd/mm/yy")
    page.insert_text((250, 658), "MEDAN",    fontsize=8, fontname="helv", color=(0,0,0))
    page.insert_text((323, 660), dd,          fontsize=8, fontname="helv", color=(0,0,0))
    page.insert_text((341, 660), mm,          fontsize=8, fontname="helv", color=(0,0,0))
    page.insert_text((370, 660), yy,          fontsize=8, fontname="helv", color=(0,0,0))

    # Nama calon di kolom TTD calon — center x=(398+504)/2=451
    insert_centered(nama_display, 451, 762+PY, 7)

    # Nama wadah — center x=155, y=762+PY (koordinat dari Colab)
    insert_centered(nama_wadah, 155, 762+PY, 7)
    # Jabatan wadah — x=141+PX=144, y=774+PY=772
    page.insert_text((144, 772), jabatan_wadah, fontsize=7, fontname="helv", color=(0,0,0))

    # Nama cabang — center x=(261+367)/2=314, y=803+PY=801
    insert_centered(nama_cabang, 314, 801, 7)
    # Jabatan cabang — x=297+PX=300, y=815+PY=813
    page.insert_text((300, 813), jabatan_cabang, fontsize=7, fontname="helv", color=(0,0,0))

    # ── Checkbox & garis ─────────────────────────────────────
    def centang(rect: fitz.Rect):
        p1 = fitz.Point(rect.x0+1, rect.y0+3)
        p2 = fitz.Point(rect.x0+3, rect.y1-1)
        p3 = fitz.Point(rect.x1-1, rect.y0+1)
        page.draw_line(p1, p2, color=(0,0,0), width=1.5)
        page.draw_line(p2, p3, color=(0,0,0), width=1.5)

    for r in [
        fitz.Rect(43,467,50,474), fitz.Rect(43,487,50,494),
        fitz.Rect(43,507,50,514), fitz.Rect(43,527,50,534),
        fitz.Rect(43,575,50,582), fitz.Rect(43,595,50,602),
        fitz.Rect(43,615,50,622), fitz.Rect(43,635,50,642),
        fitz.Rect(40,855,48,862), fitz.Rect(40,870,48,877),
        fitz.Rect(40,885,48,892),
    ]:
        centang(r)

    for x1, y1, x2, y2 in [(175,860,220,860),(175,874,220,874),(175,891,220,891)]:
        page.draw_line(fitz.Point(x1,y1), fitz.Point(x2,y2), color=(0,0,0), width=1)

    # ── Gambar stempel & TTD ─────────────────────────────────
    w_sw, h_sw = 4.2*72/2.54, 4.2*72/2.54
    w_sc, h_sc = 6.8*72/2.54, 2.0*72/2.54
    w_tw, h_tw = 4.0*72/2.54, 2.5*72/2.54
    w_tc, h_tc = 6.5*72/2.54, 3.8*72/2.54

    safe_insert_img(page, STEMPEL_WADAH,  fitz.Rect(40, 655, 40+w_sw, 655+h_sw),             "stempel_wadah")
    safe_insert_img(page, TTD_WADAH,      fitz.Rect(163-w_tw/2, 695, 163-w_tw/2+w_tw, 695+h_tw), "ttd_wadah")
    safe_insert_img(page, STEMPEL_CABANG, fitz.Rect(304-w_sc/2, 735, 304-w_sc/2+w_sc, 735+h_sc), "stempel_cabang")
    safe_insert_img(page, TTD_CABANG,     fitz.Rect(314-w_tc/2, 710, 314-w_tc/2+w_tc, 710+h_tc), "ttd_cabang")
    # TTD calon — koordinat dari Colab: fitz.Rect(390, 690, 525, 762)
    safe_insert_img(page, ttd_path,       fitz.Rect(390, 690, 525, 762),                       "ttd_calon")
    # Materai — koordinat dari Colab: fitz.Rect(235, 524, 319, 608)
    safe_insert_img(page, materai_path,   fitz.Rect(235, 524, 319, 608),                       "materai")

    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()

# ── PDF Generator: PKS ────────────────────────────────────────────────────────

def docx_to_pdf(docx_path: Path, output_pdf: Path) -> bool:
    try:
        from docx2pdf import convert
        convert(str(docx_path), str(output_pdf))
        return output_pdf.exists()
    except Exception as e:
        logger.error(f"docx2pdf error: {e}")
        raise HTTPException(500, f"Gagal konversi DOCX ke PDF: {e}")

def generate_pks(data: KeplingData, ttd_path: Optional[Path], materai_path: Optional[Path]) -> bytes:
    from docx import Document

    if not TEMPLATE_PKS.exists():
        raise FileNotFoundError("Template PKS tidak ditemukan di assets/")

    now        = datetime.datetime.now()
    hari_list  = ["Minggu","Senin","Selasa","Rabu","Kamis","Jumat","Sabtu"]
    bulan_list = ["","Januari","Februari","Maret","April","Mei","Juni",
                  "Juli","Agustus","September","Oktober","November","Desember"]
    romawi     = ["","I","II","III","IV","V","VI","VII","VIII","IX","X","XI","XII"]

    def terbilang(n: int) -> str:
        sat = ["","Satu","Dua","Tiga","Empat","Lima","Enam","Tujuh","Delapan","Sembilan","Sepuluh","Sebelas"]
        if n < 12: return sat[n]
        if n < 20: return sat[n-10] + " Belas"
        if n < 100: return sat[n//10] + " Puluh " + (sat[n%10] if n%10 else "")
        if 2000 <= n < 2100: return "Dua Ribu " + (terbilang(n-2000) if n > 2000 else "")
        return str(n)

    nama_wadah, jabatan_wadah, nama_cabang, jabatan_cabang = resolve_penandatangan(data)
    nama_display = data.namaCalonPerisai or data.namaLengkap
    no_surat     = data.nomorSurat or 1

    replacement = {
        "[NO_SURAT]":         str(no_surat),
        "[BLN_ROMAWI]":       romawi[now.month],
        "[TAHUN]":            str(now.year),
        "[HARI]":             hari_list[now.isoweekday() % 7],
        "[TGL_HURUF]":        terbilang(now.day).strip(),
        "[BLN_NAMA]":         bulan_list[now.month],
        "[THN_HURUF]":        terbilang(now.year).strip(),
        "[TGL_ANGKA]":        f"({now.strftime('%d - %m - %Y')})",
        "[NAMA_CALON_JUDUL]": nama_display.upper(),
        "[NAMA_CALON]":       nama_display.title(),
        "[KTP_CALON]":        data.noKTP,
        "[ALAMAT_CALON]":     data.alamat.title(),
        "[TELP_CALON]":       data.noTelp,
        "[EMAIL_CALON]":      data.email,
        "[NAMA_WADAH]":       nama_wadah,
        "[JABATAN_WADAH]":    jabatan_wadah,
        "[NAMA_CABANG]":      nama_cabang,
        "[JABATAN_CABANG]":   jabatan_cabang,
    }

    doc = Document(str(TEMPLATE_PKS))

    def merge_and_replace(para):
        runs = para.runs
        i = 0
        while i < len(runs):
            if "[" in runs[i].text and "]" not in runs[i].text:
                j = i + 1
                while j < len(runs):
                    runs[i].text += runs[j].text
                    runs[j].text = ""
                    if "]" in runs[i].text: break
                    j += 1
            i += 1
        for key, val in replacement.items():
            if key in para.text:
                for run in para.runs:
                    if key in run.text:
                        run.text = run.text.replace(key, val)

    all_paras = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                all_paras.extend(cell.paragraphs)
    for p in all_paras:
        merge_and_replace(p)

    with tempfile.NamedTemporaryFile(suffix=".docx", dir=TEMP_DIR, delete=False) as f:
        temp_docx = Path(f.name)
    doc.save(str(temp_docx))

    pdf_path = TEMP_DIR / (temp_docx.stem + ".pdf")
    docx_to_pdf(temp_docx, pdf_path)
    temp_docx.unlink(missing_ok=True)

    if not pdf_path.exists():
        raise HTTPException(500, "Gagal konversi DOCX ke PDF")

    pdf_doc   = fitz.open(str(pdf_path))
    last_page = pdf_doc[-1]
    safe_insert_img(last_page, materai_path, fitz.Rect(235, 494, 319, 578), "materai_pks")
    safe_insert_img(last_page, ttd_path,     fitz.Rect(145, 490, 280, 580), "ttd_calon_pks")

    buf = io.BytesIO()
    pdf_doc.save(buf)
    pdf_doc.close()
    pdf_path.unlink(missing_ok=True)
    return buf.getvalue()

# ── PDF Generator: Exam ───────────────────────────────────────────────────────

def generate_exam(data: ExamData) -> bytes:
    if not TEMPLATE_UJIAN.exists():
        raise FileNotFoundError("Template ujian tidak ditemukan di assets/")

    doc  = fitz.open(str(TEMPLATE_UJIAN))
    hal1 = doc[0]
    hal2 = doc[1]

    hal1.insert_text((187, 226), data.namaLengkap.title(), fontsize=11, fontname="hebo", color=(0,0,0))
    hal2.insert_text((103, 593), str(int(data.skor)),       fontsize=14, fontname="hebo", color=(0,0,0))

    koor = {
        1:  {"A":(90,310,0),"B":(90,323,0),"C":(90,351,0),"D":(90,363,0)},
        2:  {"A":(90,405,0),"B":(90,419,0),"C":(90,433,0),"D":(90,459,0)},
        3:  {"A":(90,515,0),"B":(90,527,0),"C":(90,555,0),"D":(90,569,0)},
        4:  {"A":(90,610,0),"B":(90,623,0),"C":(90,637,0),"D":(90,650,0)},
        5:  {"A":(90,692,0),"B":(90,705,0),"C":(90,719,0),"D":(90,733,0)},
        6:  {"A":(90,121,1),"B":(90,134,1),"C":(90,149,1),"D":(90,161,1)},
        7:  {"A":(90,217,1),"B":(90,230,1),"C":(90,244,1),"D":(90,257,1)},
        8:  {"A":(90,298,1),"B":(90,311,1),"C":(90,326,1),"D":(90,338,1)},
        9:  {"A":(90,394,1),"B":(90,406,1),"C":(90,421,1),"D":(90,433,1)},
        10: {"A":(93,476,1),"B":(93,488,1),"C":(93,517,1),"D":(93,529,1)},
    }

    def silang(page, x, y):
        sz = 4
        page.draw_line(fitz.Point(x-sz,y-sz), fitz.Point(x+sz,y+sz), color=(0,0,0), width=1.3)
        page.draw_line(fitz.Point(x-sz,y+sz), fitz.Point(x+sz,y-sz), color=(0,0,0), width=1.3)

    for i in range(1, 11):
        j = data.jawaban.get(f"Soal {i}", "").strip()
        if j: j = j[0].upper()
        if i in koor and j in koor[i]:
            x, y, hi = koor[i][j]
            silang(doc[hi], x, y)

    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()

# ── Sheets: ambil data exam ──────────────────────────────────────────────────

async def get_exam_data_for(nama: str, sheets_id: str, api_key: str) -> Optional[ExamData]:
    url  = f"https://sheets.googleapis.com/v4/spreadsheets/{sheets_id}/values/Form%20Responses%201!A:L?key={api_key}"
    async with httpx.AsyncClient(timeout=15) as client:
        resp = await client.get(url)
        if resp.status_code != 200: return None
        rows = resp.json().get("values", [])

    nama_lower = nama.lower().strip()
    for row in rows[1:]:
        if len(row) < 12: continue
        if nama_lower in str(row[11]).strip().lower() or str(row[11]).strip().lower() in nama_lower:
            skor    = float(str(row[1]).replace(",", ".")) if row[1] else 0.0
            jawaban = {f"Soal {i+1}": (str(row[2+i]).strip()[0].upper() if 2+i < len(row) and row[2+i] else "") for i in range(10)}
            return ExamData(namaLengkap=str(row[11]).strip(), skor=skor, jawaban=jawaban)
    return None

# ── Config ────────────────────────────────────────────────────────────────────

SHEETS_ID      = os.getenv("VITE_SHEETS_ID", "")
API_KEY        = os.getenv("VITE_SHEETS_API_KEY", "")
GDRIVE_API_KEY = os.getenv("VITE_SHEETS_API_KEY", "")

# ── Endpoints ─────────────────────────────────────────────────────────────────

@app.post("/generate/formulir")
async def endpoint_formulir(data: KeplingData):
    kode_pos = await resolve_kodepos(data)

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        ttd_path, materai_path = None, None

        if data.ttdUrl:
            dest = tmp / "ttd_calon.png"
            if await download_file_from_url(data.ttdUrl, dest, "ttd_calon"):
                ttd_path = dest

        if data.materaiUrl:
            dest = tmp / "materai.png"
            if await download_file_from_url(data.materaiUrl, dest, "materai"):
                materai_path = dest

        logger.info(f"[formulir] ttd={ttd_path}, materai={materai_path}, kodepos={kode_pos}")

        try:
            pdf_bytes = generate_formulir(data, ttd_path, materai_path, kode_pos)
        except FileNotFoundError as e:
            raise HTTPException(404, str(e))

    nama_file = (data.namaCalonPerisai or data.namaLengkap).replace(" ", "_").upper()
    return StreamingResponse(
        io.BytesIO(pdf_bytes), media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="PERISAI_FORM_{nama_file}.pdf"'},
    )


@app.post("/generate/pks")
async def endpoint_pks(data: KeplingData):
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        ttd_path, materai_path = None, None

        if data.ttdUrl:
            dest = tmp / "ttd_calon.png"
            if await download_file_from_url(data.ttdUrl, dest, "ttd_calon_pks"):
                ttd_path = dest

        if data.materaiUrl:
            dest = tmp / "materai.png"
            if await download_file_from_url(data.materaiUrl, dest, "materai_pks"):
                materai_path = dest

        logger.info(f"[pks] ttd={ttd_path}, materai={materai_path}")

        try:
            pdf_bytes = generate_pks(data, ttd_path, materai_path)
        except FileNotFoundError as e:
            raise HTTPException(404, str(e))

    nama_file = (data.namaCalonPerisai or data.namaLengkap).replace(" ", "_").upper()
    return StreamingResponse(
        io.BytesIO(pdf_bytes), media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="PKS_{nama_file}.pdf"'},
    )


@app.post("/generate/exam")
async def endpoint_exam(payload: dict):
    nama = payload.get("namaCalonPerisai") or payload.get("namaKepling", "")
    if not nama: raise HTTPException(400, "namaCalonPerisai wajib diisi")
    exam_data = await get_exam_data_for(nama, SHEETS_ID, API_KEY)
    if not exam_data: raise HTTPException(404, f"Data ujian untuk '{nama}' tidak ditemukan")
    try:
        pdf_bytes = generate_exam(exam_data)
    except FileNotFoundError as e:
        raise HTTPException(404, str(e))
    nama_file = nama.replace(" ", "_").upper()
    return StreamingResponse(
        io.BytesIO(pdf_bytes), media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="EXAM_{nama_file}.pdf"'},
    )


class MergeData(BaseModel):
    dokumenUrl:       str   # URL langsung file dokumen.pdf
    namaCalonPerisai: str


@app.post("/generate/merge")
async def endpoint_merge(data: MergeData):
    """Gabungkan: Form template + dokumen.pdf calon + Template Ujian."""
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)

        dokumen_path = tmp / "dokumen.pdf"
        if not await download_file_from_url(data.dokumenUrl, dokumen_path, "dokumen"):
            raise HTTPException(404, "Gagal mengunduh dokumen.pdf dari URL yang diberikan")

        parts: list[fitz.Document] = []
        if TEMPLATE_FORM.exists():   parts.append(fitz.open(str(TEMPLATE_FORM)))
        if dokumen_path.exists():    parts.append(fitz.open(str(dokumen_path)))
        if TEMPLATE_UJIAN.exists():  parts.append(fitz.open(str(TEMPLATE_UJIAN)))

        if not parts:
            raise HTTPException(500, "Tidak ada berkas yang bisa digabungkan")

        merged = fitz.open()
        for part in parts:
            merged.insert_pdf(part)
            part.close()

        buf = io.BytesIO()
        merged.save(buf)
        merged.close()

    nama_file = data.namaCalonPerisai.replace(" ", "_").upper()
    return StreamingResponse(
        io.BytesIO(buf.getvalue()), media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="COMPLETE_FILES_{nama_file}.pdf"'},
    )


@app.get("/kodepos")
async def endpoint_kodepos(kelurahan: str, kecamatan: str = ""):
    return {"kodePos": await lookup_kodepos(kelurahan, kecamatan)}


@app.get("/health")
def health():
    return {
        "status": "ok",
        "assets": {
            "template_pks":   TEMPLATE_PKS.exists(),
            "template_form":  TEMPLATE_FORM.exists(),
            "template_ujian": TEMPLATE_UJIAN.exists(),
            "stempel_wadah":  STEMPEL_WADAH.exists(),
            "ttd_wadah":      TTD_WADAH.exists(),
            "stempel_cabang": STEMPEL_CABANG.exists(),
            "ttd_cabang":     TTD_CABANG.exists(),
        },
    }